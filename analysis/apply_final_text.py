#!/usr/bin/env python
"""Rebuild manuscript.docx and figures.docx from the author's final merged text
(verbatim; Unicode superscripts kept as-is), with the rendered figures embedded.
Preserves the template's Normal style by clearing its paragraphs and re-adding.
Bold for headings and run-in leads; ("IMG", path, width_in) embeds a centred picture.
Usage: apply_final_text.py <template_dir> <out_dir> <figures_dir>
"""
import sys, docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

TPL, OUT = sys.argv[1], sys.argv[2]
FIG = sys.argv[3] if len(sys.argv) > 3 else None
IMG = {"fig1": f"{FIG}/fig1/fig1.png", "ed1": f"{FIG}/ed1/ed1.png",
       "ed2": f"{FIG}/ed2/ed2.png", "ed3": f"{FIG}/ed3/ed3.png"} if FIG else {}

# ("H", text)=full-bold heading ; ("B", text)=body ; ("L", lead, rest)=bold lead + body
SUP = {}  # superscripts are provided inline as Unicode glyphs (¹²³…)

ABSTRACT = ("Kondepudi et al. evaluated how well NeuroVFM perceives brain scans by having it write a "
    "report and then asking a separate AI system to judge urgency from that report. Because the judge "
    "performed well on radiologists’ reports, they reasoned that errors reflected what NeuroVFM saw. "
    "Using the publicly released model on a public set of head CT scans, we found that this reasoning "
    "misses a step. NeuroVFM’s own diagnostic classifier, which reads the same image features, "
    "identified urgent scans as well as the full report pipeline did, but the two missed different "
    "scans. In most urgent scans the pipeline missed, the written report did not mention the "
    "abnormality and usually described the scan as normal, even though the classifier often scored "
    "that abnormality highly. Report-based evaluation therefore measures perception and report "
    "writing together, and it may understate what the model perceives.")

MAIN = [
"As foundation models have grown, clinical AI systems are increasingly assembled as chains: one "
"model perceives, a second writes about what it perceived, and a third reads that text to make a "
"decision. Each added step brings capability, but it also adds a place where information can be "
"lost, and an evaluation of the whole chain can attribute an error to the wrong step. Kondepudi et "
"al.¹ introduced NeuroVFM, a neuroimaging foundation model trained on routine clinical scans, and "
"made a persuasive case that such models outperform general-purpose frontier models. Its frozen "
"encoder supports two readouts, a diagnostic classifier that scores 82 CT diagnoses and a language "
"model that writes findings. Yet perception and triage were evaluated only through that prose, "
"which GPT-5 then read to assign acuity. Because GPT-5 assigned acuity accurately from radiologists’ "
"findings, the authors concluded that acuity differences “reflect perception quality rather than "
"reasoning ability”¹, and prospective errors were attributed to missed radiographic findings rather "
"than clinician mis-triage¹. This design isolates the judge, but not the encoder from the report "
"writer. A finding can be perceived by the encoder yet omitted when the report is written, and "
"because both readouts are public, we tested whether this occurs.",

"The classifier is a small attention-pooled network² of about 0.95 million parameters, and the "
"report path uses a 14-billion-parameter language model³,⁴ whose findings GPT-5 reads to assign "
"acuity (Fig. 1a). We left the model unchanged and evaluated both readouts side by side. For the "
"classifier, we scored each study by its highest score across the critical findings, with no "
"language model and no additional training. We ran both readouts on CQ500⁵, a public set of 472 "
"non-contrast head CTs with consensus labels from three radiologists, 206 of which contained at "
"least one of nine critical findings. We used the released weights, prompts and GPT-5 screening "
"model, and both readouts received identical series and CT windows. To compare them fairly, we set "
"the classifier’s threshold so that it flagged the same fraction of studies as the report pipeline, "
"a choice that uses no labels, and applied that threshold throughout (Methods).",

"Used directly as a triage score, the classifier performed as well as the report pipeline. It "
"separated urgent from non-urgent studies with an area under the receiver operating characteristic "
"curve of 0.949 (95% CI 0.928–0.967, Fig. 1b). When set to flag the same fraction of studies as the "
"report pipeline (34.1%), it detected the same fraction of urgent studies (73.3% for both, "
"difference 0.000, 95% CI −0.029 to 0.030). We do not claim that the classifier is the better "
"triager. The two readouts did not, however, miss the same studies. Of 206 urgent studies, 134 were "
"caught by both and 38 by neither, and the remaining 34 split evenly, with 17 caught only by the "
"classifier and 17 only by the report. Because both readouts draw on the same encoder, a study "
"missed by the report is not necessarily a study the encoder failed to see.",

"The 55 studies missed by the report pipeline make this concrete. In 5, the report named a critical "
"finding but GPT-5 did not escalate it. In the other 50, the report named none of the study’s "
"critical findings, and most read as explicitly normal, for example, “Study is unremarkable.” These "
"were complete reports rather than truncated or empty outputs, they were reproduced verbatim on "
"regeneration, and the language model had received the same images as the classifier, including the "
"bone window. In 40 of the 50, the classifier scored the missed finding above 0.5 (median 0.92), "
"and 13 of the 50 would have been flagged by the classifier at the matched alarm rate (Fig. 1c). "
"The omitted findings included calvarial fracture, intraparenchymal hemorrhage, subdural hematoma "
"and mass effect (Extended Data Fig. 1). The original evaluation would count all 50 as failures of "
"perception. For many of them, the encoder had represented the finding, and it was lost when the "
"report was written. A report that reads as normal is more harmful than a missing flag, because it "
"reassures and moves the scan down the worklist.",

"Several checks support this reading. The language model did not simply call everything normal, "
"describing 96% of non-urgent studies as normal but only 24% of urgent ones. The results were "
"similar with a different screening model (Claude) and with a second judge of whether each finding "
"was mentioned (agreement κ = 0.97). Because the authors’ screening rules do not escalate some "
"findings¹,⁶, we repeated the analysis without the four studies whose only critical finding was a "
"calvarial fracture, and the two readouts still detected the same fraction of urgent studies "
"(difference 0.000, 95% CI −0.030 to 0.030). CQ500 also provides no clinical indication, which the "
"report model normally receives. In the authors’ own experiments, withholding the indication "
"lowered NeuroVFM’s acuity accuracy only slightly, from about 0.80 to 0.76 (Extended Data Fig. 9d "
"of ref. 1). When we supplied a generic indication, 38 of the 40 silent reports with high "
"classifier scores remained silent under “head trauma” and under “headache”. The report pipeline’s "
"sensitivity did not improve (0.728 and 0.694, compared with 0.733 without an indication), whereas "
"the classifier’s was unchanged at 0.733 (Extended Data Fig. 2).",

"Our analysis has limits. CQ500 contains more urgent studies than routine practice (44%, compared "
"with 13% in the authors’ prospective week), the classifier’s scores are not calibrated to data "
"from other institutions⁷, and we could not examine the 21 prospective misses themselves.",

"None of this questions the strength of NeuroVFM. If anything, its encoder perceives more than the "
"report-based evaluation gave it credit for. The authors envision reasoning models acting on "
"structured, calibrated outputs from NeuroVFM¹, and the diagnostic classifier already produces "
"exactly such outputs. We suggest that triage, and any evaluation of perception, be measured "
"through those outputs as well as through generated text. The most direct test is in the authors’ "
"hands: classifier scores for the 21 prospective urgent misses would show how many were failures of "
"perception and how many were failures of report writing. More broadly, when a chain of models is "
"judged only by the text it produces, errors of writing can be mistaken for errors of seeing⁸, and "
"developers may then build larger models when what they need is a different readout.",
]

REFS = [
"1. Kondepudi, A. et al. Health system learning enables generalist neuroimaging models. Nat. Med. 32, 2831–2837 (2026).",
"2. Ilse, M., Tomczak, J. & Welling, M. Attention-based deep multiple instance learning. In Proc. 35th International Conference on Machine Learning 2127–2136 (PMLR, 2018).",
"3. Liu, H., Li, C., Li, Y. & Lee, Y. J. Improved baselines with visual instruction tuning. In 2024 IEEE/CVF Conference on Computer Vision and Pattern Recognition 26286–26296 (IEEE, 2024).",
"4. Yang, A. et al. Qwen3 technical report. Preprint at https://arxiv.org/abs/2505.09388 (2025).",
"5. Chilamkurthy, S. et al. Deep learning algorithms for detection of critical findings in head CT scans: a retrospective study. Lancet 392, 2388–2396 (2018).",
"6. Babiarz, L. S. et al. Neuroradiology critical findings lists: survey of neuroradiology training programs. AJNR Am. J. Neuroradiol. 34, 735–739 (2013).",
"7. Guo, C., Pleiss, G., Sun, Y. & Weinberger, K. Q. On calibration of modern neural networks. In Proc. 34th International Conference on Machine Learning 1321–1330 (PMLR, 2017).",
"8. Manrai, A. K. Medical AI has a measurement problem. Nature 655, 1138–1139 (2026).",
"9. Efron, B. & Tibshirani, R. J. An Introduction to the Bootstrap (Chapman & Hall, 1993).",
]

LEGENDS = [
("Fig. 1.", " NeuroVFM’s diagnostic classifier triages as well as its written reports, and many "
"reports omit findings the classifier scored highly. a, The released NeuroVFM system and the change "
"we evaluated. One frozen encoder feeds two readouts that do not share information. In the original "
"study, perception and triage were evaluated through the report path, in which a 14-billion-"
"parameter language model writes findings that GPT-5 reads to assign acuity (orange). We also read "
"the diagnostic classifier (about 0.95 million parameters) directly (blue). Shown for one CQ500 "
"study with subarachnoid hemorrhage, the classifier scored the finding at 0.99 and flagged the "
"study, whereas the generated report read “Study is unremarkable” and the study was not escalated. "
"b, Sensitivity for urgent studies versus fraction of studies flagged (n = 472, 206 urgent). Blue "
"curve and band, the classifier’s highest score across the nine critical findings used directly as "
"a triage score, with 95% CI (AUROC 0.949, 95% CI 0.928–0.967). Orange point, the report pipeline "
"(sensitivity 0.733 at a flag rate of 0.341). Blue point, the classifier at the same flag rate, a "
"threshold chosen without labels. The difference in sensitivity was 0.000 (95% CI −0.029 to 0.030). "
"c, Classifier score for the missed finding in each of the 55 urgent studies the report pipeline "
"missed. Filled blue, report silent and study flagged by the classifier at the matched alarm rate "
"(n = 13). Light blue, report silent and score above 0.5 but study not flagged (n = 27). Grey, "
"report silent and score of 0.5 or below (n = 10). Open orange, the report named a critical finding "
"that GPT-5 did not escalate (n = 5). Dotted line, score of 0.5. Tick, median score of silent "
"reports scored above 0.5 (0.92). Confidence intervals in b are from 2,000 study-level bootstrap "
"resamples."),
("Extended Data Fig. 1.", " Classification of report-pipeline misses under two definitions and two "
"screening models. Each miss was classified as reasoning (the report named a critical finding that "
"was not escalated), decoding (the report was silent and the classifier detected the finding) or "
"perception (the report was silent and the classifier did not detect the finding). a, Classification "
"of the 55 misses when classifier detection is defined as a score above 0.5 on the highest-scoring "
"positive finding (40 decoding, 10 perception, 5 reasoning), and when it is defined as flagging the "
"study at the matched alarm rate (13 decoding, 37 perception, 5 reasoning). b, The same "
"classification with Claude as the screening model (of 48 misses, 34 decoding, 10 perception and 4 "
"reasoning with the 0.5 score, and 11 decoding, 33 perception and 4 reasoning at the matched alarm "
"rate). Acuity agreement between the two screeners was 96.6%. c, Decoding misses by critical finding "
"with the 0.5 score (calvarial fracture 9, intraparenchymal hemorrhage 8, intracranial hemorrhage "
"8, subdural hematoma 5, mass effect 5, subarachnoid hemorrhage 3, epidural hematoma 2). d, Fraction "
"of generated reports reading as normal among urgent (24%) and non-urgent (96%) studies. e, "
"Agreement on whether each reference finding was mentioned, between GPT-5 and Claude as judges (mean "
"κ 0.97) and between each judge and a rule-based detector (κ 0.92 or higher for hemorrhage "
"findings)."),
("Extended Data Fig. 2.", " Clinical indication, urgency definition, calibration and prevalence. a, "
"Classification of misses with no clinical indication (as in the main analysis, 40 decoding, 10 "
"perception and 5 reasoning of 55), with the indication “head trauma” (40, 10 and 6 of 56) and with "
"the indication “headache” (47, 11 and 5 of 63). Of the 40 decoding misses in the main analysis, 38 "
"persisted under each indication and 37 under both. b, Results after excluding the four studies "
"whose only critical finding was a calvarial fracture (n = 468, 202 urgent). The sensitivity "
"difference remained 0.000 (95% CI −0.030 to 0.030, AUROC 0.948), and the 51 misses comprised 37 "
"decoding, 9 perception and 5 reasoning misses. c, Reliability diagram for classifier scores pooled "
"over the nine critical findings (expected calibration error 0.25). The classifier discriminates "
"well on external data but is not calibrated to it, so we describe its outputs as scores rather than "
"probabilities. d, Sensitivity versus fraction flagged with the flag-rate axis reweighted to the "
"13% urgent prevalence of the original prospective study."),
("Extended Data Fig. 3.", " Representative decoding misses. Six CQ500 studies in which the generated "
"report omitted an urgent finding. For each, a key axial slice, the classifier’s scores for the "
"nine critical findings and the generated report verbatim. The examples span subarachnoid "
"hemorrhage, intraparenchymal hemorrhage, subdural hematoma, intracranial hemorrhage, calvarial "
"fracture and mass effect. In each, the classifier scored the finding above 0.96 while the report "
"read as normal."),
]

METHODS = [
("Cohort and reference standard.", " CQ500 is a public set of non-contrast head CT studies with "
"independent reads from three radiologists⁵. Of 491 studies with reads, 472 had loadable images and "
"a consensus label and formed the analysis set. A finding was positive when at least two of three "
"readers marked it. Nine findings were treated as critical: intracranial, intraparenchymal, "
"intraventricular, epidural and subarachnoid hemorrhage, subdural hematoma, calvarial fracture, "
"mass effect and midline shift. A study was urgent if any was present (206 of 472, 44%). Kondepudi "
"et al. evaluated encoders on CQ500 by training probes on CQ500 itself, and in that setting NeuroVFM "
"underperformed a report-supervised model on most tasks¹. We instead applied the released "
"classifier, trained on the authors’ institutional data, without any CQ500 training, and evaluated "
"a study-level urgency score rather than individual labels, so the two analyses are not directly "
"comparable. Against the authors’ consensus critical-findings list¹,⁶, the 82 NeuroVFM CT labels "
"covered all critical findings on the list that can be assessed on head CT (16 of its 40 entries, "
"including hemorrhage, acute stroke, herniation, hydrocephalus, mass effect, edema, aneurysm, "
"arteriovenous malformation, venous sinus thrombosis, vascular dissection and occlusion, skull "
"fracture, pneumocephalus, abscess and foreign body). The remaining entries are spinal or soft-"
"tissue findings outside the scope of a non-contrast head CT."),
("Model structure and readouts.", " We used the released NeuroVFM encoder, CT diagnostic classifier "
"(mlinslab/neurovfm-dx-ct, an attention-pooled gated multilayer perceptron, 768→384→82, 952,264 "
"parameters) and report model (mlinslab/neurovfm-llm) without modification. Both readouts received "
"the same pooled diagnostic series for each study through the same preprocessor (1 × 1 × 4 mm "
"resampling, with brain, blood and bone windows and no cap on series). Reports were generated with "
"the released settings (beam search with four beams, no sampling, maximum 512 new tokens) and an "
"empty clinical indication, matching the image-only input of the classifier. Each CQ500 finding was "
"mapped before analysis to one or more of the 82 classifier labels, and the score for a finding was "
"the maximum over its mapped labels (Supplementary Table 1)."),
("Screening and mention judgment.", " Generated reports were screened for acuity with GPT-5 "
"(gpt-5-2025-08-07) using the authors’ critical-findings prompt¹. GPT-5 also judged whether each "
"reference finding was mentioned in each report. Claude Sonnet 4.5 (claude-sonnet-4-5) served as a "
"second screener and judge for sensitivity analyses. Mention judgments were compared between models "
"(Cohen’s κ) and against a rule-based detector."),
("Direct readout and matched threshold.", " In the direct readout, each study was scored by its "
"maximum classifier score across the nine critical findings. The threshold was set so that the "
"classifier flagged the same fraction of studies as the report pipeline. This uses only the "
"distribution of scores and no labels, so the comparison carries no optimism from threshold tuning. "
"The same threshold defined classifier detection in the two-by-two comparison of readouts."),
("Classification of misses.", " Each urgent study was classified as detected or missed by each "
"readout. A study missed by the report pipeline was a reasoning miss if its report named at least "
"one of the study’s positive critical findings. Otherwise, it was a decoding miss if the classifier "
"scored the highest-scoring positive finding above 0.5, and a perception miss if that score was 0.5 "
"or below. Requiring the report to be silent on all positive findings is conservative. Scoring the "
"highest-scoring positive finding, however, favors a decoding classification, so we also classified "
"misses using flags at the matched alarm rate, which does not depend on choosing a finding (13 "
"decoding, 37 perception and 5 reasoning misses, Extended Data Fig. 1). Every decoding-miss report "
"was inspected for empty, truncated or errored output."),
("Urgency definition.", " The authors’ screening rules exempt some findings from escalation¹. CQ500 "
"does not record lesion size, so we could not apply size-based exemptions. We therefore repeated "
"all analyses after excluding studies whose only critical finding was a calvarial fracture, the "
"finding most likely to be exempt when isolated."),
("Indication ablation and determinism.", " Reports for all 472 studies were regenerated with the "
"clinical indication “head trauma” and, separately, “headache”. Persistence of decoding misses "
"under both indications would exclude absent indications as their cause. Resolution of many would "
"indicate that the sensitivity of the report path depends on prompt context, whereas the "
"classifier’s does not. Reports for all decoding misses were regenerated at the released batch size "
"of one study per pass to test determinism. All 40 reproduced verbatim, confirming that the "
"omissions are deterministic under the model’s beam-search decoding rather than sampling artifacts."),
("Statistics.", " Confidence intervals were obtained from 2,000 bootstrap resamples of studies⁹, "
"with the matched threshold re-selected within each resample. Calibration was summarized by expected "
"calibration error over score deciles⁷. Analyses used Python 3.10, PyTorch 2.5.0 (CUDA 12.4), "
"FlashAttention 2.6.3, NumPy and SciPy."),
("Use of large language models.", " GPT-5 (OpenAI) and Claude (Anthropic) were used as screening "
"models and mention judges as described above. Claude was also used to assist with code "
"development, figure generation and manuscript editing. The authors reviewed and verified all "
"output and take full responsibility for the work."),
("Ethics.", " This study used only publicly available, de-identified data."),
("Reporting summary.", " Further information on research design is available in the Nature Portfolio "
"Reporting Summary linked to this article."),
("Data availability.", " CQ500 is publicly available at http://headctstudy.qure.ai/dataset. "
"Study-level classifier scores, generated reports, screening outputs and mention judgments "
"underlying Fig. 1 and Extended Data Figs. 1–3 are provided as Source Data and at "
"https://github.com/Jonah-Keller/neurovfm-matters-arising."),
("Code availability.", " Code for inference, label mapping, classification of misses and statistical "
"analysis is available at https://github.com/Jonah-Keller/neurovfm-matters-arising. The NeuroVFM "
"code and weights are available from the original authors at "
"https://github.com/MLNeurosurg/neurovfm."),
]

def manuscript_items():
    it = []
    it += [("H", "Matters Arising"),
           ("H", "Report-based evaluation of NeuroVFM does not isolate visual perception"),
           ("H", "Arising from: A. Kondepudi et al. Nature Medicine https://doi.org/10.1038/s41591-026-04497-1 (2026)"),
           ("H", "Authors: Jonah N. Keller¹, Visish M. Srinivasan¹*"),
           ("H", "Affiliations:"),
           ("B", "¹Department of Neurosurgery, Perelman School of Medicine, University of Pennsylvania, Philadelphia, Pennsylvania"),
           ("B", "* Corresponding author:"),
           ("B", "visishs@upenn.edu"),
           ("H", "Abstract"), ("B", ABSTRACT),
           ("H", "Main Text")]
    it += [("B", p) for p in MAIN]
    it += [("H", "References")] + [("B", r) for r in REFS]
    it += [("H", "Figure Legends")]
    keys = ["fig1", "ed1", "ed2", "ed3"]
    for k, (a, b) in zip(keys, LEGENDS):
        if k in IMG: it.append(("IMG", IMG[k], 6.5))
        it.append(("L", a, b))
    it += [("H", "Methods")] + [("L", a, b) for a, b in METHODS]
    it += [("H", "Acknowledgements"),
           ("B", "We thank the authors of ref. 1 for releasing the NeuroVFM code and weights, and the creators of CQ500 for making their data public."),
           ("H", "Author contributions"),
           ("B", "J.N.K. and V.M.S. conceived the study. J.N.K. performed the analyses and drafted the manuscript. V.M.S. supervised the work. Both authors revised and approved the final manuscript."),
           ("H", "Competing interests"),
           ("B", "The authors declare no competing interests.")]
    return it

def figures_items():
    it = [("H", "Figures")]
    keys = ["fig1", "ed1", "ed2", "ed3"]
    labels = ["[Insert Figure 1 here]", "[Insert Extended Data Figure 1 here]",
              "[Insert Extended Data Figure 2 here]", "[Insert Extended Data Figure 3 here]"]
    for k, lab, (lead, rest) in zip(keys, labels, LEGENDS):
        it.append(("IMG", IMG[k], 6.5) if k in IMG else ("B", lab))
        it += [("L", lead, rest), ("B", "")]
    return it

def build(template, out, items, double=True):
    d = docx.Document(template)
    for p in list(d.paragraphs):
        p._element.getparent().remove(p._element)
    for item in items:
        if item[0] == "IMG":
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format; pf.line_spacing = 1.0
            pf.space_before = Pt(6); pf.space_after = Pt(6)
            p.add_run().add_picture(item[1], width=Inches(item[2]))
            continue
        p = d.add_paragraph()
        if item[0] == "H":
            r = p.add_run(item[1]); r.bold = True
        elif item[0] == "B":
            p.add_run(item[1])
        elif item[0] == "L":
            rb = p.add_run(item[1]); rb.bold = True
            p.add_run(item[2])
        pf = p.paragraph_format
        pf.line_spacing = 2.0 if double else 1.0        # double-spaced manuscript
        pf.space_after = Pt(0)
    d.save(out)
    wc = sum(len((i[1] + (i[2] if i[0] == "L" else "")).split())
             for i in items if i[0] in ("H", "B", "L"))
    return wc

mwc = build(f"{TPL}/NeuroVFM_Matters_Arising_manuscript.docx",
            f"{OUT}/NeuroVFM_Matters_Arising_manuscript.docx", manuscript_items())
build(f"{TPL}/NeuroVFM_Matters_Arising_figures.docx",
      f"{OUT}/NeuroVFM_Matters_Arising_figures.docx", figures_items())
# main-text word count (Main Text paragraphs only)
main_wc = sum(len(p.split()) for p in MAIN)
print(f"saved manuscript + figures. Main-Text words={main_wc}; total doc words≈{mwc}")
